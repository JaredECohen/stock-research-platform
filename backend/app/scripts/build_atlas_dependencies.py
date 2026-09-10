"""Build ``data/industry_knowledge/atlas_dependencies.json`` — the checked-in
cross-industry dependency graph the Portfolio Manager's cross-industry
snapshot reads (FEAT-003).

Usage (from ``backend/``)::

    python -m app.scripts.build_atlas_dependencies          # rebuild the JSON
    python -m app.scripts.build_atlas_dependencies --check  # exit 1 if stale

Two kinds of link, two provenances, one file:

* ``edges`` — the Atlas workbook's ``Dependencies`` sheet
  (``docs/research/Investment_Research_Atlas.xlsx``): ~27 analyst causal
  hypotheses of the form origin → destination with a transmission
  mechanism, example exposures (symbols), an invalidation and the next
  evidence to look for. Reading the workbook needs ``openpyxl``, which is
  a dev-interpreter package and deliberately NOT a backend dependency;
  this script is therefore dev-time only and the JSON it writes is
  committed. Without ``openpyxl`` the script prints why and exits 0 — it
  never writes a file with the edges silently missing.
* ``relationships`` — the universe map's ``cross_industry_relationships``
  (themes with explicit 8-digit sub-industry codes, a monitor and a
  ``failure_of_inference``), read from the merged knowledge JSON that
  ``build_industry_knowledge`` writes so this script never parses the map
  itself (the "one taxonomy source" rule). Always emitted.

Both are resolved to industry-group codes so the snapshot can label a
spillover with its origin and destination groups: an edge's example
exposures go through the knowledge JSON's ``security_reference`` (symbol →
sub-industry codes; unresolvable labels such as ``"regional banks"`` are
listed, not dropped), a relationship's codes are already sub-industries.
Every record carries its ``source`` so a consumer can say where a link
came from — an Atlas edge is a hypothesis (its ``status`` says so), not a
measured correlation.

Output is byte-deterministic for a given workbook + knowledge JSON, so
``--check`` doubles as a drift test.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parents[2]
DEFAULT_ATLAS = REPO_ROOT / "docs" / "research" / "Investment_Research_Atlas.xlsx"
DEFAULT_KNOWLEDGE = SCRIPT_DIR.parent / "data" / "industry_knowledge" / "gics_industries_2026.json"
DEFAULT_OUTPUT = SCRIPT_DIR.parent / "data" / "industry_knowledge" / "atlas_dependencies.json"

DEPENDENCIES_SHEET = "Dependencies"
SETTINGS_SHEET = "Settings"
# Header cells of the Dependencies sheet, in order, and the keys they map to.
EDGE_HEADER = (
    "Edge Id", "Origin", "Destination", "Transmission",
    "Example Exposures", "Invalidation", "Next Evidence", "Status",
)
EDGE_KEYS = (
    "edge_id", "origin", "destination", "transmission",
    "example_exposures", "invalidation", "next_evidence", "status",
)
SOURCE_ATLAS = "atlas_dependencies_sheet"
SOURCE_MAP = "universe_map_cross_industry_relationships"


def _die(message: str) -> None:
    print(f"build_atlas_dependencies: {message}", file=sys.stderr)
    raise SystemExit(1)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _relative(path: Path) -> str:
    try:
        return path.resolve().relative_to(REPO_ROOT).as_posix()
    except ValueError:
        return path.name


# --- the workbook ------------------------------------------------------------


def openpyxl_available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def read_edges(atlas: Path) -> tuple[list[dict[str, Any]], str | None]:
    """The Dependencies sheet as ``[{edge_id, origin, …}]`` plus the
    workbook's snapshot date. Raises ``ImportError`` without openpyxl."""
    import openpyxl

    if not atlas.is_file():
        _die(f"atlas workbook not found: {atlas}")
    wb = openpyxl.load_workbook(atlas, read_only=True, data_only=True)
    if DEPENDENCIES_SHEET not in wb.sheetnames:
        _die(f"{atlas} has no '{DEPENDENCIES_SHEET}' sheet")

    rows = list(wb[DEPENDENCIES_SHEET].iter_rows(values_only=True))
    header_at = next(
        (i for i, row in enumerate(rows) if row and _text(row[0]) == EDGE_HEADER[0]), None,
    )
    if header_at is None:
        _die(f"'{DEPENDENCIES_SHEET}' sheet has no '{EDGE_HEADER[0]}' header row")
    header = tuple(_text(c) for c in rows[header_at][: len(EDGE_HEADER)])
    if header != EDGE_HEADER:
        _die(f"'{DEPENDENCIES_SHEET}' header changed: {header} (expected {EDGE_HEADER})")

    edges: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rows[header_at + 1:]:
        cells = [_text(c) for c in (row or ())] + [""] * len(EDGE_KEYS)
        edge_id = cells[0]
        if not edge_id:
            continue  # trailing blank rows
        if not (edge_id.startswith("D") and edge_id[1:].isdigit()):
            _die(f"edge id {edge_id!r} is not of the form D<n>")
        if edge_id in seen:
            _die(f"duplicate edge id {edge_id}")
        seen.add(edge_id)
        record = dict(zip(EDGE_KEYS, cells[: len(EDGE_KEYS)]))
        for key in ("origin", "destination", "transmission"):
            if not record[key]:
                _die(f"edge {edge_id} has an empty '{key}'")
        record["example_exposures"] = [
            part.strip() for part in record["example_exposures"].split(";") if part.strip()
        ]
        edges.append(record)
    if not edges:
        _die(f"'{DEPENDENCIES_SHEET}' sheet holds no edges")

    snapshot: str | None = None
    if SETTINGS_SHEET in wb.sheetnames:
        for row in wb[SETTINGS_SHEET].iter_rows(values_only=True):
            if row and _text(row[0]).lower() == "snapshot date" and len(row) > 1:
                value = row[1]
                if isinstance(value, (datetime, date)):
                    snapshot = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                elif value:
                    snapshot = _text(value)[:10]
                break
    wb.close()
    return edges, snapshot


# --- the knowledge JSON ------------------------------------------------------


def load_knowledge(path: Path) -> dict[str, Any]:
    if not path.is_file():
        _die(f"knowledge JSON not found: {path} — run build_industry_knowledge first")
    try:
        doc = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        _die(f"{path} is not valid JSON: {exc}")
    for key in ("sectors", "cross_industry_relationships", "security_reference"):
        if key not in doc:
            _die(f"{path} lacks '{key}' — it predates the sub-industry layer; rebuild it")
    return doc


def active_sub_industry_codes(doc: dict[str, Any]) -> set[str]:
    return {
        sub["code"]
        for sector in doc["sectors"]
        for group in sector["industry_groups"]
        for industry in group["industries"]
        for sub in industry.get("sub_industries", [])
    }


def _group_codes(codes: list[str]) -> list[str]:
    return sorted({c[:4] for c in codes})


def resolve_exposures(
    exposures: list[str], security_reference: dict[str, dict[str, Any]],
) -> tuple[dict[str, list[str]], list[str]]:
    """``({SYMBOL: [code8…]}, [labels that are not a known symbol])``.

    Exposure cells mix symbols (``NVDA``), foreign listings
    (``HPS.A@CA``) and prose (``regional banks``); only an exact
    upper-cased symbol in the security reference resolves. The rest is
    reported so nobody mistakes "no code" for "no exposure".
    """
    resolved: dict[str, list[str]] = {}
    unresolved: list[str] = []
    for label in exposures:
        entry = security_reference.get(label.upper())
        if entry is None:
            unresolved.append(label)
        else:
            resolved[label.upper()] = list(entry.get("codes", []))
    return resolved, unresolved


def build_payload(
    knowledge: Path = DEFAULT_KNOWLEDGE,
    atlas: Path = DEFAULT_ATLAS,
) -> dict[str, Any]:
    """Assemble the JSON document. Requires openpyxl (see ``main`` for the
    graceful no-openpyxl exit)."""
    doc = load_knowledge(knowledge)
    active = active_sub_industry_codes(doc)
    if not active:
        _die(f"{knowledge} carries no active sub-industries")
    security_reference = doc["security_reference"]

    raw_edges, snapshot = read_edges(atlas)
    edges: list[dict[str, Any]] = []
    for raw in raw_edges:
        resolved, unresolved = resolve_exposures(raw["example_exposures"], security_reference)
        codes = sorted({c for entry in resolved.values() for c in entry})
        unknown = [c for c in codes if c not in active]
        if unknown:
            _die(f"edge {raw['edge_id']} resolves to codes not in the knowledge JSON: {unknown}")
        edges.append({
            **raw,
            "source": SOURCE_ATLAS,
            "resolved_exposures": resolved,
            "unresolved_exposures": unresolved,
            "sub_industry_codes": codes,
            "industry_group_codes": _group_codes(codes),
        })

    relationships: list[dict[str, Any]] = []
    for rel in doc["cross_industry_relationships"]:
        codes = list(rel.get("codes", []))
        unknown = [c for c in codes if c not in active]
        if unknown:
            _die(f"relationship {rel.get('theme')!r} references unknown codes {unknown}")
        relationships.append({
            "theme": rel["theme"],
            "mechanism": rel["mechanism"],
            "codes": codes,
            "industry_group_codes": _group_codes(codes),
            "monitor": rel.get("monitor", ""),
            "failure_of_inference": rel.get("failure_of_inference", ""),
            "source_ids": list(rel.get("source_ids", [])),
            "source": SOURCE_MAP,
        })

    return {
        "_doc": (
            f"Provenance: generated from {_relative(atlas)} (sheet '{DEPENDENCIES_SHEET}') "
            f"and {_relative(knowledge)} by app/scripts/build_atlas_dependencies.py; do not "
            "edit by hand — rerun `python -m app.scripts.build_atlas_dependencies` from "
            "backend/ (needs openpyxl in the dev interpreter). `edges` are the Atlas's "
            "analyst causal hypotheses (economic exposure graph, not estimated "
            "correlations); `relationships` are the universe map's themed cross-industry "
            "links. Codes are the public GICS numeric hierarchy; exposures resolve through "
            "the map's economic security reference, which is research examples, not "
            "licensed issuer GICS assignments."
        ),
        "generated_from": {"atlas": _relative(atlas), "knowledge": _relative(knowledge)},
        "atlas_sha256": hashlib.sha256(atlas.read_bytes()).hexdigest(),
        "atlas_snapshot_date": snapshot,
        "knowledge_taxonomy_version": doc.get("taxonomy_version"),
        "knowledge_source_sha256": doc.get("source_sha256"),
        "knowledge_map_source_sha256": doc.get("map_source_sha256"),
        "knowledge_map_as_of": doc.get("map_as_of"),
        "edge_source": SOURCE_ATLAS,
        "relationship_source": SOURCE_MAP,
        "edge_count": len(edges),
        "relationship_count": len(relationships),
        "edges": edges,
        "relationships": relationships,
    }


def render(payload: dict[str, Any]) -> bytes:
    return (json.dumps(payload, indent=2, ensure_ascii=False) + "\n").encode("utf-8")


def build(
    knowledge: Path = DEFAULT_KNOWLEDGE, atlas: Path = DEFAULT_ATLAS, output: Path = DEFAULT_OUTPUT,
) -> dict[str, Any]:
    payload = build_payload(knowledge, atlas)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(render(payload))
    return payload


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Compile the Atlas dependency edges and the map's cross-industry "
                    "relationships into one JSON.",
    )
    parser.add_argument("--atlas", type=Path, default=DEFAULT_ATLAS,
                        help=f"Atlas workbook (default: {DEFAULT_ATLAS})")
    parser.add_argument("--knowledge", type=Path, default=DEFAULT_KNOWLEDGE,
                        help=f"merged knowledge JSON (default: {DEFAULT_KNOWLEDGE})")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT,
                        help=f"JSON to write (default: {DEFAULT_OUTPUT})")
    parser.add_argument("--check", action="store_true",
                        help="Write nothing; exit 1 if the output is missing or differs "
                             "from what the sources produce.")
    args = parser.parse_args(argv)

    if not openpyxl_available():
        # Not an error: the backend image has no openpyxl by design. The
        # committed JSON stays as it is; say so rather than fail a CI job
        # or, worse, write a file with no edges.
        print(
            "build_atlas_dependencies: openpyxl is not importable in this interpreter; "
            "the Atlas workbook cannot be read. Nothing written — run this from the dev "
            f"interpreter to refresh {args.output.name}.",
        )
        return 0

    payload = build_payload(args.knowledge, args.atlas)
    rendered = render(payload)
    summary = f"{payload['edge_count']} edges / {payload['relationship_count']} relationships"

    if args.check:
        if not args.output.is_file():
            print(f"{args.output} is missing ({summary} would be written)", file=sys.stderr)
            return 1
        if args.output.read_bytes() != rendered:
            print(
                f"{args.output} is stale — rerun `python -m app.scripts.build_atlas_dependencies`",
                file=sys.stderr,
            )
            return 1
        print(f"{args.output} is up to date ({summary})")
        return 0

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_bytes(rendered)
    print(f"wrote {args.output} ({summary})")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
